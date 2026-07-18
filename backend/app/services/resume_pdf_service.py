from __future__ import annotations

import re
import uuid
from datetime import date
from pathlib import Path
from uuid import UUID

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import ActingCredit, ActorProfile, Asset
from app.services.file_storage_service import FileStorageService


GENERATED_RESUME_TAG = "generated-resume"
GENERATED_RESUME_DOCX_TAG = "generated-resume-docx"


class ResumePdfService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.storage = FileStorageService()
        self.output_dir = get_settings().upload_dir / "generated_resumes"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def regenerate_for_actor(self, actor_profile_id: UUID) -> Asset:
        actor = self.db.get(ActorProfile, actor_profile_id)
        if not actor:
            raise ValueError("Actor profile not found")

        credits = list(
            self.db.scalars(
                select(ActingCredit)
                .where(ActingCredit.actor_profile_id == actor_profile_id)
                .order_by(
                    ActingCredit.section_order.asc(),
                    ActingCredit.category.asc(),
                    ActingCredit.display_order.asc(),
                    ActingCredit.created_at.asc(),
                )
            )
        )
        generated_id = uuid.uuid4()
        path = self.output_dir / f"{actor_profile_id}-{generated_id}.pdf"
        docx_path = self.output_dir / f"{actor_profile_id}-{generated_id}.docx"
        self._write_pdf(path, actor, credits)
        self._write_docx(docx_path, actor, credits)
        self._delete_previous_generated_resumes(actor_profile_id)
        today = date.today()
        asset = Asset(
            actor_profile_id=actor_profile_id,
            asset_name="Generated Acting Resume",
            asset_type="Resume",
            local_file_path=str(path),
            original_filename=f"{self._slug(actor.name)}-acting-resume.pdf",
            mime_type="application/pdf",
            file_size_bytes=path.stat().st_size,
            description="Auto-generated from the structured Acting Resume Builder.",
            tags=[GENERATED_RESUME_TAG, "acting-resume", "auto-generated"],
            archetype_names=[],
            ai_suggested_tags=[],
            ai_suggested_archetypes=[],
            analysis_status="complete",
            analysis_explanation="Generated from structured acting credits.",
            upload_date=today,
            last_updated_date=today,
            expiration_warning_date=today,
            freshness_status="Current",
        )
        self.db.add(asset)
        docx_asset = Asset(
            actor_profile_id=actor_profile_id,
            asset_name="Generated Acting Resume DOCX",
            asset_type="Resume",
            local_file_path=str(docx_path),
            original_filename=f"{self._slug(actor.name)}-acting-resume.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size_bytes=docx_path.stat().st_size,
            description="Editable Word resume generated from the structured Acting Resume Builder.",
            tags=[GENERATED_RESUME_DOCX_TAG, "acting-resume", "auto-generated", "editable"],
            archetype_names=[],
            ai_suggested_tags=[],
            ai_suggested_archetypes=[],
            analysis_status="complete",
            analysis_explanation="Generated editable document from structured acting credits.",
            upload_date=today,
            last_updated_date=today,
            expiration_warning_date=today,
            freshness_status="Current",
        )
        self.db.add(docx_asset)
        return asset

    def current_generated_resume(self, actor_profile_id: UUID) -> Asset | None:
        return self.db.scalars(
            select(Asset)
            .where(
                Asset.actor_profile_id == actor_profile_id,
                Asset.asset_type == "Resume",
                Asset.tags.any(GENERATED_RESUME_TAG),
            )
            .order_by(Asset.updated_at.desc())
        ).first()

    def current_generated_resume_docx(self, actor_profile_id: UUID) -> Asset | None:
        return self.db.scalars(
            select(Asset)
            .where(
                Asset.actor_profile_id == actor_profile_id,
                Asset.asset_type == "Resume",
                Asset.tags.any(GENERATED_RESUME_DOCX_TAG),
            )
            .order_by(Asset.updated_at.desc())
        ).first()

    def _delete_previous_generated_resumes(self, actor_profile_id: UUID) -> None:
        assets = list(
            self.db.scalars(
                select(Asset).where(
                    Asset.actor_profile_id == actor_profile_id,
                    Asset.asset_type == "Resume",
                    Asset.tags.overlap([GENERATED_RESUME_TAG, GENERATED_RESUME_DOCX_TAG]),
                )
            )
        )
        for asset in assets:
            path = asset.local_file_path
            self.db.delete(asset)
            self.storage.delete_file(path)

    def _write_pdf(self, path: Path, actor: ActorProfile, credits: list[ActingCredit]) -> None:
        page_width, page_height = letter
        pdf = canvas.Canvas(str(path), pagesize=letter)
        left = 1.0 * inch
        bottom = 0.5 * inch
        y = page_height - 0.58 * inch
        col_1 = left
        col_2 = 3.85 * inch
        col_3 = 5.75 * inch

        def set_font(name: str, size: int) -> None:
            pdf.setFont(name, size)

        def draw(text: str, x: float, yy: float, font: str = "Times-Roman", size: int = 11) -> None:
            set_font(font, size)
            pdf.drawString(x, yy, self._clean_inline(text))

        def draw_underlined(text: str, yy: float) -> float:
            draw(text, left, yy, "Times-Bold", 13)
            width = stringWidth(text, "Times-Bold", 13)
            pdf.line(left, yy - 1.5, left + width, yy - 1.5)
            return yy - 15

        def new_page_if_needed(yy: float, required: float = 42) -> float:
            if yy < bottom + required:
                pdf.showPage()
                return page_height - 0.55 * inch
            return yy

        draw(actor.name, left, y, "Times-Bold", 20)
        y -= 24
        draw(actor.union_status or actor.sag_status or "", left, y, "Times-Roman", 16)
        y -= 22
        representation = self._primary_representation(actor)
        if representation:
            y = draw_underlined("Agent:", y)
            if representation.agent_name or representation.agency_name:
                agent_line = " - ".join(part for part in [representation.agent_name, representation.agency_name] if part)
                draw(agent_line, left, y, "Times-Roman", 12)
                y -= 14
            contact_line = "  -  ".join(part for part in [representation.agent_email, representation.agent_phone] if part)
            if contact_line:
                draw(contact_line, left, y, "Times-Roman", 12)
                y -= 18
        y -= 10

        for section_title, items in self._resume_sections(credits):
            if not items:
                continue
            y = new_page_if_needed(y, 55)
            y = draw_underlined(section_title, y)
            for credit in items:
                y = new_page_if_needed(y)
                first, second, third = self._resume_line_values(section_title, credit)
                draw(first, col_1, y)
                if second:
                    draw(second, col_2, y)
                if third:
                    draw(third, col_3, y)
                y -= 13
            y -= 11

        skills = self._skill_summary(actor, credits)
        if skills:
            y = new_page_if_needed(y, 70)
            y = draw_underlined("Special Skills", y)
            set_font("Times-Roman", 11)
            for line in self._wrap_text(skills, "Times-Roman", 11, page_width - left - 0.5 * inch):
                y = new_page_if_needed(y)
                pdf.drawString(left, y, line)
                y -= 13

        pdf.save()

    def _write_docx(self, path: Path, actor: ActorProfile, credits: list[ActingCredit]) -> None:
        from docx import Document
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.shared import Inches, Pt

        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.5)

        styles = document.styles
        styles["Normal"].font.name = "Times New Roman"
        styles["Normal"].font.size = Pt(11)

        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(0)
        run = title.add_run(actor.name)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(20)

        meta = document.add_paragraph()
        meta.paragraph_format.space_after = Pt(10)
        meta_run = meta.add_run(actor.union_status or actor.sag_status or "")
        meta_run.font.name = "Times New Roman"
        meta_run.font.size = Pt(16)

        representation = self._primary_representation(actor)
        if representation:
            agent_label = document.add_paragraph()
            agent_label.paragraph_format.space_after = Pt(0)
            agent_run = agent_label.add_run("Agent:")
            agent_run.bold = True
            agent_run.underline = True
            agent_run.font.name = "Times New Roman"
            agent_run.font.size = Pt(12)

            if representation.agent_name or representation.agency_name:
                agent_line = document.add_paragraph()
                agent_line.paragraph_format.space_after = Pt(0)
                text = " - ".join(part for part in [representation.agent_name, representation.agency_name] if part)
                run = agent_line.add_run(text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

            contact_line = "  -  ".join(part for part in [representation.agent_email, representation.agent_phone] if part)
            if contact_line:
                contact = document.add_paragraph()
                contact.paragraph_format.space_after = Pt(12)
                run = contact.add_run(contact_line)
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

        for section_title, items in self._resume_sections(credits):
            if not items:
                continue
            self._docx_heading(document, section_title)
            for credit in items:
                first, second, third = self._resume_line_values(section_title, credit)
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                tabs = paragraph.paragraph_format.tab_stops
                tabs.add_tab_stop(Inches(2.75), WD_TAB_ALIGNMENT.LEFT)
                tabs.add_tab_stop(Inches(4.75), WD_TAB_ALIGNMENT.LEFT)
                run = paragraph.add_run(f"{first}\t{second}\t{third}".rstrip())
                run.font.name = "Times New Roman"
                run.font.size = Pt(10.5)

        skills = self._skill_summary(actor, credits)
        if skills:
            self._docx_heading(document, "Special Skills")
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            run = paragraph.add_run(skills)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
        document.save(path)

    def _docx_heading(self, document, text: str) -> None:
        from docx.shared import Pt

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = True
        run.underline = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)

    def _resume_line_values(self, section_title: str, credit: ActingCredit) -> tuple[str, str, str]:
        if section_title == "Training":
            instructor = credit.instructor or ""
            institution = credit.institution or ""
            teacher = ", ".join(part for part in [instructor, institution] if part)
            return (credit.class_or_program or "", "", teacher or credit.year or "")
        if section_title == "Special Skills":
            return (credit.skill_name or "", "", "")
        third = credit.network_or_distributor or credit.production_company or credit.director or ""
        if section_title == "Theater":
            third = credit.production_company or credit.network_or_distributor or credit.director or ""
        return (
            credit.project_title or "",
            credit.role_or_character or credit.role_type or "",
            third,
        )

    def _resume_sections(self, credits: list[ActingCredit]) -> list[tuple[str, list[ActingCredit]]]:
        enabled = [credit for credit in credits if credit.section_enabled and credit.category != "Special Skills"]
        categories = {
            "Television": ["Television"],
            "Film/New Media": ["Film", "New Media"],
            "Commercial": ["Commercial", "Industrial", "Print", "Voiceover"],
            "Theater": ["Theater"],
            "Training": ["Training"],
            "Other": ["Other"],
        }
        sections: list[tuple[str, list[ActingCredit]]] = []
        for title, source_categories in categories.items():
            items = [
                credit
                for credit in enabled
                if credit.category in source_categories and self._has_resume_content(credit)
            ]
            if items:
                sections.append((title, sorted(items, key=lambda item: (item.display_order, item.created_at))))
        return sections

    def _has_resume_content(self, credit: ActingCredit) -> bool:
        return bool(
            credit.project_title
            or credit.role_or_character
            or credit.role_type
            or credit.production_company
            or credit.network_or_distributor
            or credit.director
            or credit.class_or_program
            or credit.instructor
            or credit.institution
        )

    def _primary_representation(self, actor: ActorProfile):
        active = [representation for representation in actor.representations if representation.active]
        return active[0] if active else None

    def _skill_summary(self, actor: ActorProfile, credits: list[ActingCredit]) -> str:
        skills: list[str] = []
        accents: list[str] = []
        for credit in credits:
            if credit.section_enabled and credit.category == "Special Skills" and credit.skill_name:
                if self._looks_like_accent(credit.skill_name):
                    accents.append(self._clean_accent_name(credit.skill_name))
                else:
                    skills.append(credit.skill_name)
        skills.extend(actor.skills or [])
        accents.extend(actor.accents or [])
        clean_skills = self._dedupe_phrases(skills)
        clean_accents = self._dedupe_phrases([self._clean_accent_name(accent) for accent in accents])
        parts = clean_skills
        if clean_accents:
            parts.append(f"Dialects: {', '.join(clean_accents)}")
        return ", ".join(parts)

    def _dedupe_phrases(self, values: list[str]) -> list[str]:
        seen: set[str] = set()
        result: list[str] = []
        for value in values:
            cleaned = self._clean_inline(value)
            if not cleaned:
                continue
            key = re.sub(r"[^a-z0-9]+", "", cleaned.lower())
            if key in seen:
                continue
            seen.add(key)
            result.append(cleaned)
        return result

    def _looks_like_accent(self, value: str) -> bool:
        return bool(re.search(r"\b(accent|dialect|rp|southern|new york|boston|west african|kenyan|valley girl)\b", value, flags=re.I))

    def _clean_accent_name(self, value: str) -> str:
        cleaned = self._clean_inline(value)
        cleaned = re.sub(r"^dialects?:\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s+accent$", "", cleaned, flags=re.I)
        return cleaned

    def _wrap_text(self, text: str, font: str, size: int, max_width: float) -> list[str]:
        lines: list[str] = []
        current = ""
        for word in text.split():
            candidate = f"{current} {word}".strip()
            if stringWidth(candidate, font, size) <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
        return lines

    def _clean_inline(self, value: str | None) -> str:
        return re.sub(r"\s+", " ", value or "").strip()

    def _slug(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "actor"
